"""Materialize Copilot-edited design documents as trusted Canonical changes."""

from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
from dataclasses import dataclass
from nturl2path import url2pathname as windows_url2pathname
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlsplit

from openpyxl import load_workbook
from psycopg import Connection

from operamind.application.document_diff import (
    DocumentDiffService,
    DocumentSnapshotBuildResult,
)
from operamind.contracts import ContractCatalog
from operamind.domain import (
    CanonicalDocumentNodeBuilder,
    CanonicalSnapshot,
    ChangeConfidence,
    ChangeReviewStatus,
    StructuredChangeBuilder,
)
from operamind.domain.canonical_facts import normalize_business_value
from operamind.domain.document_conventions import DocumentConvention
from operamind.infrastructure.documents import DocumentSignalExtractorRegistry
from operamind.infrastructure.postgres import (
    ArtifactRepository,
    CanonicalDocumentSlice,
    CanonicalRepository,
    DocumentNodeRepository,
    DocumentSnapshotWrite,
    ProfileRepository,
    SnapshotStatus,
)
from operamind.profiles import ProfileCatalog


@dataclass(frozen=True, slots=True)
class _PreparedDocument:
    source: CanonicalDocumentSlice
    path: Path
    digest: str
    target_document_version_id: str
    target: DocumentSnapshotBuildResult
    convention: DocumentConvention


@dataclass(frozen=True, slots=True)
class DocumentFieldEdit:
    """One Copilot-requested Canonical field update applied by OperaMind."""

    document_id: str
    stable_key: str
    field: str
    new_value: str

    def __post_init__(self) -> None:
        values = (self.document_id, self.stable_key, self.field, self.new_value)
        if any(not value.strip() for value in values):
            raise ValueError("Document field edit values must not be blank")
        if len(self.new_value) > 20_000:
            raise ValueError("Document field edit value exceeds 20,000 characters")
        if self.new_value.lstrip().startswith("="):
            raise ValueError("Document field edit must not introduce a spreadsheet formula")


@dataclass(frozen=True, slots=True)
class _StagedDocument:
    source: CanonicalDocumentSlice
    path: Path
    replacement: Path
    backup: Path


@dataclass(frozen=True, slots=True)
class CopilotDocumentChangeResult:
    source_snapshot_id: str
    target_snapshot_id: str
    document_ids: tuple[str, ...]
    source_paths: tuple[Path, ...]
    change_refs: tuple[str, ...]

    def to_dict(self) -> dict[str, object]:
        return {
            "source_snapshot_id": self.source_snapshot_id,
            "target_snapshot_id": self.target_snapshot_id,
            "document_ids": list(self.document_ids),
            "source_paths": [str(path) for path in self.source_paths],
            "change_refs": list(self.change_refs),
        }


class CopilotDocumentChangeService:
    """Compare modified trusted source files to their persisted Canonical memberships."""

    def __init__(self, *, connection: Connection[Any], repository_root: Path) -> None:
        root = repository_root.resolve()
        self._connection = connection
        self._contracts = ContractCatalog.load(root / "contracts")
        self._profiles = ProfileCatalog.load(root / "profiles")
        self._canonical = CanonicalRepository(connection, self._contracts)
        self._profile_repository = ProfileRepository(connection, self._profiles)
        self._artifacts = ArtifactRepository(connection, self._contracts)
        self._nodes = DocumentNodeRepository(connection)
        self._node_builder = CanonicalDocumentNodeBuilder()
        self._extractors = DocumentSignalExtractorRegistry.default()
        self._document_diff = DocumentDiffService(
            extractors=self._extractors,
            contracts=self._contracts,
        )
        self._change_builder = StructuredChangeBuilder()

    def materialize(
        self,
        *,
        project_id: str,
        analysis_case_id: str,
        coding_task_id: str,
        source_snapshot_id: str,
        document_ids: tuple[str, ...],
    ) -> CopilotDocumentChangeResult:
        if not document_ids or len(document_ids) != len(set(document_ids)):
            raise ValueError("Selected design document IDs must be non-empty and unique")
        sources = tuple(
            self._required_source(
                project_id=project_id,
                source_snapshot_id=source_snapshot_id,
                document_id=document_id,
            )
            for document_id in sorted(document_ids)
        )
        source_paths = tuple(_trusted_file_path(source.source_ref) for source in sources)
        digests = tuple(_file_digest(path) for path in source_paths)
        unchanged = [
            source.document_id
            for source, digest in zip(sources, digests, strict=True)
            if digest == source.content_digest
        ]
        if unchanged:
            raise ValueError(
                "Selected design documents have no file change: " + ", ".join(unchanged)
            )
        target_snapshot_id = _id(
            "copilot-document-snapshot",
            project_id,
            coding_task_id,
            source_snapshot_id,
            *(
                f"{source.document_id}:{digest}"
                for source, digest in zip(sources, digests, strict=True)
            ),
        )
        prepared = tuple(
            self._prepare_document(
                source=source,
                path=path,
                digest=digest,
                target_snapshot_id=target_snapshot_id,
                project_id=project_id,
            )
            for source, path, digest in zip(sources, source_paths, digests, strict=True)
        )
        source_snapshot = CanonicalSnapshot(
            snapshot_id=source_snapshot_id,
            facts=tuple(fact for item in sources for fact in item.snapshot.facts),
        )
        target_snapshot = CanonicalSnapshot(
            snapshot_id=target_snapshot_id,
            facts=tuple(fact for item in prepared for fact in item.target.snapshot.facts),
        )
        source_keys = [fact.fact.stable_key for fact in source_snapshot.facts]
        target_keys = [fact.fact.stable_key for fact in target_snapshot.facts]
        if len(source_keys) != len(set(source_keys)) or len(target_keys) != len(set(target_keys)):
            raise ValueError("Selected design documents contain duplicate Canonical Stable Keys")
        changes = tuple(
            change
            for domain in sorted({item.convention.document_type for item in prepared})
            for change in self._change_builder.diff(
                project_id=project_id,
                source=CanonicalSnapshot(
                    snapshot_id=source_snapshot_id,
                    facts=tuple(
                        fact
                        for item in prepared
                        if item.convention.document_type == domain
                        for fact in item.source.snapshot.facts
                    ),
                ),
                target=CanonicalSnapshot(
                    snapshot_id=target_snapshot_id,
                    facts=tuple(
                        fact
                        for item in prepared
                        if item.convention.document_type == domain
                        for fact in item.target.snapshot.facts
                    ),
                ),
                domain=domain,
                confidence=ChangeConfidence.HIGH,
                review_status=ChangeReviewStatus.ACCEPTED,
            )
        )
        if not changes:
            raise ValueError("Modified design documents produced no semantic Canonical changes")
        for change in changes:
            self._contracts.validate_artifact(change.to_artifact())
        with self._connection.transaction():
            for item in prepared:
                self._persist_target(project_id=project_id, item=item)
            self._canonical.store_changes(changes)
            for change in changes:
                self._artifacts.store(
                    artifact_id=change.change_id,
                    project_id=project_id,
                    analysis_case_id=analysis_case_id,
                    artifact=change.to_artifact(),
                )
        return CopilotDocumentChangeResult(
            source_snapshot_id=source_snapshot_id,
            target_snapshot_id=target_snapshot_id,
            document_ids=tuple(source.document_id for source in sources),
            source_paths=source_paths,
            change_refs=tuple(change.change_id for change in changes),
        )

    def apply_and_materialize(
        self,
        *,
        project_id: str,
        analysis_case_id: str,
        coding_task_id: str,
        source_snapshot_id: str,
        document_ids: tuple[str, ...],
        document_edits: tuple[DocumentFieldEdit, ...],
    ) -> CopilotDocumentChangeResult:
        """Apply bounded Canonical XLSX field edits, then materialize their semantic diff."""

        if not document_edits:
            raise ValueError("Structured document edits must not be empty")
        if {edit.document_id for edit in document_edits} != set(document_ids):
            raise ValueError("Structured document edits must cover exactly selected documents")
        edit_keys = [
            (edit.document_id, edit.stable_key, edit.field) for edit in document_edits
        ]
        if len(edit_keys) != len(set(edit_keys)):
            raise ValueError("Structured document edits contain duplicate Canonical fields")
        sources = {
            document_id: self._required_source(
                project_id=project_id,
                source_snapshot_id=source_snapshot_id,
                document_id=document_id,
            )
            for document_id in sorted(document_ids)
        }
        staged = tuple(
            self._stage_xlsx_edits(
                source=sources[document_id],
                edits=tuple(
                    edit for edit in document_edits if edit.document_id == document_id
                ),
            )
            for document_id in sorted(document_ids)
        )
        try:
            for item in staged:
                if _file_digest(item.path) != item.source.content_digest:
                    raise ValueError(
                        "Design document changed after its structured edit was prepared: "
                        f"{item.source.document_id}"
                    )
            for item in staged:
                shutil.copy2(item.path, item.backup)
            for item in staged:
                os.replace(item.replacement, item.path)
            try:
                return self.materialize(
                    project_id=project_id,
                    analysis_case_id=analysis_case_id,
                    coding_task_id=coding_task_id,
                    source_snapshot_id=source_snapshot_id,
                    document_ids=document_ids,
                )
            except Exception:
                for item in staged:
                    if item.backup.exists():
                        os.replace(item.backup, item.path)
                raise
        finally:
            for item in staged:
                item.replacement.unlink(missing_ok=True)
                item.backup.unlink(missing_ok=True)

    def _stage_xlsx_edits(
        self,
        *,
        source: CanonicalDocumentSlice,
        edits: tuple[DocumentFieldEdit, ...],
    ) -> _StagedDocument:
        path = _trusted_file_path(source.source_ref)
        if path.suffix.casefold() != ".xlsx":
            raise ValueError(
                "Structured document editing currently supports only XLSX design documents"
            )
        if _file_digest(path) != source.content_digest:
            raise ValueError(
                "Design document differs from the discovered Canonical baseline: "
                f"{source.document_id}"
            )
        facts = {item.fact.stable_key: item.fact for item in source.snapshot.facts}
        workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
        replacement = _temporary_sibling(path, "replacement")
        backup = _temporary_sibling(path, "backup")
        try:
            for edit in edits:
                fact = facts.get(edit.stable_key)
                if fact is None:
                    raise ValueError(
                        "Structured document edit Stable Key is outside the selected document: "
                        f"{edit.stable_key}"
                    )
                baseline_value = fact.values.get(edit.field)
                evidence = [
                    item
                    for item in fact.field_evidence
                    if item.canonical_field == edit.field
                ]
                if baseline_value is None or len(evidence) != 1:
                    raise ValueError(
                        "Structured document edit field is not uniquely backed by "
                        "Canonical evidence: "
                        f"{edit.stable_key}/{edit.field}"
                    )
                source_refs = evidence[0].source_refs
                if len(source_refs) != 1:
                    raise ValueError(
                        "Structured document edit field maps to multiple source cells: "
                        f"{edit.stable_key}/{edit.field}"
                    )
                sheet_name, coordinate = _xlsx_location(path, source_refs[0])
                try:
                    cell = workbook[sheet_name][coordinate]
                except KeyError as error:
                    raise ValueError(
                        f"Structured document edit worksheet does not exist: {sheet_name}"
                    ) from error
                if cell.data_type == "f":
                    raise ValueError("Structured document edit must not replace a formula cell")
                current_value = "" if cell.value is None else str(cell.value)
                if normalize_business_value(current_value) != baseline_value:
                    raise ValueError(
                        "Structured document edit source cell differs from Canonical evidence: "
                        f"{edit.stable_key}/{edit.field}"
                    )
                cell.value = edit.new_value
            workbook.save(replacement)
            os.chmod(replacement, path.stat().st_mode)
        except Exception:
            replacement.unlink(missing_ok=True)
            backup.unlink(missing_ok=True)
            raise
        finally:
            workbook.close()
        return _StagedDocument(
            source=source,
            path=path,
            replacement=replacement,
            backup=backup,
        )

    def _required_source(
        self,
        *,
        project_id: str,
        source_snapshot_id: str,
        document_id: str,
    ) -> CanonicalDocumentSlice:
        source = self._canonical.get_document_slice(
            project_id=project_id,
            snapshot_id=source_snapshot_id,
            document_id=document_id,
        )
        if source is None:
            raise ValueError(
                "Selected design document is not in the discovered Canonical Snapshot"
            )
        return source

    def _prepare_document(
        self,
        *,
        source: CanonicalDocumentSlice,
        path: Path,
        digest: str,
        target_snapshot_id: str,
        project_id: str,
    ) -> _PreparedDocument:
        profile = self._profile_repository.get_version(source.profile_version_id)
        if profile is None or profile.get("profile_type") != "DocumentConventionProfile":
            raise ValueError("Canonical design document has no valid Convention Profile")
        convention = DocumentConvention.from_validated_profile(profile)
        fact_types = {fact.fact.fact_type for fact in source.snapshot.facts}
        if len(fact_types) != 1:
            raise ValueError(
                "Selected Canonical document contains multiple Fact types and cannot be "
                "compared without an explicit mapping"
            )
        target = self._document_diff.build_snapshot(
            path=path,
            snapshot_id=target_snapshot_id,
            fact_type=next(iter(fact_types)),
            convention=convention,
            stable_key_namespace=_source_namespace(source),
        )
        return _PreparedDocument(
            source=source,
            path=path,
            digest=digest,
            target_document_version_id=_id(
                "copilot-document-version",
                project_id,
                source.document_id,
                digest,
            ),
            target=target,
            convention=convention,
        )

    def _persist_target(self, *, project_id: str, item: _PreparedDocument) -> None:
        nodes = self._node_builder.build(
            snapshot=item.target.snapshot,
            document_version_id=item.target_document_version_id,
            logical_name=item.source.logical_name,
            document_type=item.convention.document_type,
        )
        self._canonical.store_snapshot(
            DocumentSnapshotWrite(
                project_id=project_id,
                document_id=item.source.document_id,
                document_version_id=item.target_document_version_id,
                logical_name=item.source.logical_name,
                source_ref=item.path.as_uri(),
                content_digest=item.digest,
                extractor_ref=self._extractors.extractor_ref(item.path),
                profile_version_id=item.source.profile_version_id,
                selected_variant_id=item.target.selected_variant_id,
                status=SnapshotStatus.COMMITTED,
                snapshot=item.target.snapshot,
                selected_variant_ids=item.target.selected_variant_ids,
                fact_variant_ids=item.target.fact_variant_ids,
            )
        )
        self._nodes.store_nodes(
            project_id=project_id,
            snapshot_id=item.target.snapshot.snapshot_id,
            nodes=nodes,
        )


def _temporary_sibling(path: Path, purpose: str) -> Path:
    with tempfile.NamedTemporaryFile(
        prefix=f".{path.stem}.operamind-{purpose}-",
        suffix=path.suffix,
        dir=path.parent,
        delete=False,
    ) as stream:
        return Path(stream.name)


def _xlsx_location(path: Path, source_ref: str) -> tuple[str, str]:
    filename, separator, location = source_ref.partition("#")
    sheet_name, cell_separator, coordinate = location.rpartition("!")
    if (
        separator != "#"
        or cell_separator != "!"
        or filename != path.name
        or not sheet_name.strip()
        or not coordinate.strip()
    ):
        raise ValueError("Canonical XLSX field evidence has an invalid source cell reference")
    return sheet_name, coordinate


def _trusted_file_path(source_ref: str) -> Path:
    parsed = urlsplit(source_ref)
    if parsed.scheme != "file" or parsed.netloc not in {"", "localhost"}:
        raise ValueError("Canonical design document source must be a local file URI")
    path = _file_uri_path(source_ref).resolve(strict=True)
    if not path.is_file():
        raise ValueError("Canonical design document source is not a regular file")
    return path


def _source_namespace(source: CanonicalDocumentSlice) -> str | None:
    """Preserve the namespace used by multi-document baseline Snapshots."""

    expected = f"{source.document_id}/"
    local_keys = [fact.fact.stable_key.partition(":")[2] for fact in source.snapshot.facts]
    if local_keys and all(local_key.startswith(expected) for local_key in local_keys):
        return source.document_id
    return None


def _file_uri_path(source_ref: str, *, platform_name: str | None = None) -> Path:
    """Convert a local file URI without losing a Windows drive letter."""

    parsed = urlsplit(source_ref)
    if parsed.scheme != "file" or parsed.netloc not in {"", "localhost"}:
        raise ValueError("Canonical design document source must be a local file URI")
    path = unquote(parsed.path)
    platform = os.name if platform_name is None else platform_name
    if platform == "nt":
        path = windows_url2pathname(path)
    return Path(path)


def _file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _id(prefix: str, *values: str) -> str:
    material = "\0".join(values).encode()
    return f"{prefix}-{hashlib.sha256(material).hexdigest()[:24]}"
